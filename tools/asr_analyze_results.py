#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
测试结果分析工具
用于分析result.csv或result.txt文件，统计识别率并生成报告
支持两种格式：
1. result.csv: 原有格式，包含Voice, ExpectedText, ActualText, Result等字段
2. result.txt: 新格式，包含VoiceInputPath, ExpectResults, ActualResults, isPass等字段
"""

import os
import argparse
import csv
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def is_valid_result_txt(file_path):
    """
    验证result.txt文件的第一行是否符合要求
    
    Args:
        file_path: 文件路径
        
    Returns:
        bool: 是否符合要求
    """
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            first_line = f.readline().strip()
            # 期望的第一行格式
            expected_header = "Solution	Scene	Suite	case_index	VoiceInputPath	Preconditions	ExpectResults	ActualResults	isPass"
            return first_line == expected_header
    except Exception:
        return False


def find_all_result_files(root_dir):
    """
    查找所有result.csv或result.txt文件
    优先查找result.csv，如果找不到则查找result.txt（需要验证第一行格式）
    
    Args:
        root_dir: 根目录路径
        
    Returns:
        结果文件路径列表，每个元素是 (file_path, file_type) 元组
        file_type: 'csv' 或 'txt'
    """
    result_files = []
    root_path = Path(root_dir)
    
    # 优先递归查找所有result.csv文件
    csv_files = list(root_path.rglob('result.csv'))
    if csv_files:
        for csv_file in csv_files:
            result_files.append((str(csv_file), 'csv'))
        return sorted(result_files, key=lambda x: x[0])
    
    # 如果找不到csv文件，查找result.txt文件
    txt_files = list(root_path.rglob('result.txt'))
    for txt_file in txt_files:
        if is_valid_result_txt(str(txt_file)):
            result_files.append((str(txt_file), 'txt'))
    
    return sorted(result_files, key=lambda x: x[0])


def extract_prefix_value(text, prefix):
    """
    提取带前缀的值，例如从 "asr:打开主驾车窗" 中提取 "打开主驾车窗"
    
    Args:
        text: 原始文本
        prefix: 前缀（如 "asr:" 或 "voice:"）
        
    Returns:
        提取后的值，如果没有前缀则返回原文本
    """
    if not text:
        return text
    text = text.strip()
    if text.startswith(prefix):
        return text[len(prefix):].strip()
    return text


def convert_txt_to_csv_format(txt_data):
    """
    将result.txt格式的数据转换为result.csv格式的数据结构
    
    Args:
        txt_data: result.txt格式的数据列表
        
    Returns:
        转换为csv格式的数据列表
    """
    csv_data = []
    
    for row in txt_data:
        # 提取VoiceInputPath（去掉voice:前缀）
        voice_input = row.get('VoiceInputPath', '').strip()
        voice = extract_prefix_value(voice_input, 'voice:')
        
        # 提取ExpectResults（去掉asr:前缀）
        expect_results = row.get('ExpectResults', '').strip()
        expected_text = extract_prefix_value(expect_results, 'asr:')
        
        # 提取ActualResults（去掉asr:前缀）
        actual_results = row.get('ActualResults', '').strip()
        actual_text = extract_prefix_value(actual_results, 'asr:')
        
        # 转换isPass字段（True/False -> PASS/FAIL）
        is_pass = row.get('isPass', '').strip()
        if is_pass.lower() == 'true':
            result = 'PASS'
        elif is_pass.lower() == 'false':
            result = 'FAIL'
        else:
            result = 'FAIL'  # 默认FAIL
        
        # 构建csv格式的数据字典
        csv_row = {
            'Voice': voice,
            'ExpectedText': expected_text,
            'ActualText': actual_text,
            'Result': result,
            'Lang': '',  # txt格式中没有这些字段
            'Confidence': '',
            'Grade': ''
        }
        
        csv_data.append(csv_row)
    
    return csv_data


def read_result_file(file_path, file_type):
    """
    读取结果文件内容（支持csv和txt格式）
    
    Args:
        file_path: 文件路径
        file_type: 文件类型，'csv' 或 'txt'
        
    Returns:
        数据列表，每个元素是一个字典（统一为csv格式的数据结构）
    """
    data = []
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            if file_type == 'csv':
                # 读取csv格式
                reader = csv.DictReader(f, delimiter='\t')
                for row in reader:
                    data.append(row)
            elif file_type == 'txt':
                # 读取txt格式
                reader = csv.DictReader(f, delimiter='\t')
                txt_data = []
                for row in reader:
                    txt_data.append(row)
                # 转换为csv格式
                data = convert_txt_to_csv_format(txt_data)
    except Exception as e:
        print(f"读取文件 {file_path} 时出错: {e}")
    
    return data


def is_text_match(expected, actual):
    """
    判断预期文本和实际文本是否匹配（兼容大小写）
    如果两个结果的.lower()相等，就算匹配
    
    Args:
        expected: 预期文本
        actual: 实际文本
        
    Returns:
        bool: 是否匹配
    """
    if not expected or not actual:
        return False
    return expected.strip().lower() == actual.strip().lower()


def analyze_scenario_results(data):
    """
    分析场景结果
    每个ExpectedText对应多个音频，只要有一个PASS就算正确，全部FAIL才算错误
    兼容大小写：如果ExpectedText和ActualText的.lower()相等，就算正确
    
    Args:
        data: CSV数据列表
        
    Returns:
        scenario_stats: 场景统计信息字典
        failed_scenarios: 全错场景的详细信息列表
    """
    # 按ExpectedText分组
    scenario_groups = defaultdict(list)
    
    for row in data:
        expected_text = row.get('ExpectedText', '').strip()
        if expected_text:  # 忽略空值
            scenario_groups[expected_text].append(row)
    
    scenario_stats = {}
    failed_scenarios = []
    
    for expected_text, rows in scenario_groups.items():
        # 统计PASS和FAIL数量，同时检查大小写兼容的情况
        pass_count = 0
        fail_count = 0
        corrected_count = 0  # 通过大小写兼容纠正的数量
        
        for r in rows:
            original_result = r.get('Result', '').strip()
            expected = r.get('ExpectedText', '').strip()
            actual = r.get('ActualText', '').strip()
            
            # 如果原始结果是PASS，直接计数
            if original_result == 'PASS':
                pass_count += 1
            # 如果原始结果是FAIL，但通过大小写判断是匹配的，也算PASS
            elif original_result == 'FAIL' and is_text_match(expected, actual):
                pass_count += 1
                corrected_count += 1
                # 更新Result字段为PASS（纠正）
                r['Result'] = 'PASS'
                r['Corrected'] = True  # 标记为已纠正
            else:
                fail_count += 1
        
        total_count = len(rows)
        
        # 判断场景是否正确：至少有一个PASS就算正确
        is_correct = pass_count > 0
        
        scenario_stats[expected_text] = {
            'total': total_count,
            'pass': pass_count,
            'fail': fail_count,
            'is_correct': is_correct,
            'corrected': corrected_count  # 通过大小写兼容纠正的数量
        }
        
        # 如果全错，记录详细信息
        if not is_correct:
            failed_audio_details = []
            for row in rows:
                failed_audio_details.append({
                    'voice': row.get('Voice', '').strip(),
                    'expected': row.get('ExpectedText', '').strip(),
                    'actual': row.get('ActualText', '').strip() or 'None',
                    'result': row.get('Result', '').strip(),
                    'confidence': row.get('Confidence', '').strip() or 'None'
                })
            
            failed_scenarios.append({
                'expected_text': expected_text,
                'total_audio': total_count,
                'audio_details': failed_audio_details
            })
    
    return scenario_stats, failed_scenarios


def get_scenario_name(file_path):
    """
    从文件路径中提取场景名称
    
    Args:
        file_path: 文件路径
        
    Returns:
        场景名称
    """
    path_parts = Path(file_path).parts
    scenario_name = "未知场景"
    for i, part in enumerate(path_parts):
        if part == 'solution_filter' and i > 0:
            scenario_name = path_parts[i - 1]
            break
    # 如果没找到，尝试从路径中提取（相对路径的情况）
    if scenario_name == "未知场景":
        abs_path = os.path.abspath(file_path)
        if 'BiGuoji' in abs_path:
            parts = abs_path.split('BiGuoji')
            if len(parts) > 1:
                remaining = parts[1].lstrip('/')
                if remaining:
                    scenario_name = remaining.split('/')[0] or "未知场景"
    # 通用兜底：从路径中提取 父目录_子目录 形式，如 24MM中文本地必过功能集_发话一览_1
    if scenario_name == "未知场景" and len(path_parts) >= 3:
        # 路径格式: .../功能集名/序号/result.txt，提取 "功能集名_序号"
        parent_dir = path_parts[-3]  # 如 24MM中文本地必过功能集_发话一览
        sub_dir = path_parts[-2]    # 如 1
        scenario_name = f"{parent_dir}_{sub_dir}"
    return scenario_name


def fill_detail_sheet(detail_ws, failed_scenarios):
    """
    填充详细错误报告sheet的数据
    
    Args:
        detail_ws: worksheet对象
        failed_scenarios: 全错场景的详细信息列表
    """
    detail_headers = ['序号', '预期识别结果', '音频数量', '音频文件路径', '实际识别结果', '识别状态', '置信度']
    detail_columns = {'A': 8, 'B': 30, 'C': 12, 'D': 60, 'E': 30, 'F': 12, 'G': 15}
    
    for idx, header in enumerate(detail_headers, 1):
        detail_ws.cell(row=1, column=idx, value=header)
    
    setup_header_style(detail_ws, 1, detail_columns)
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    row_num = 2
    for idx, failed_scenario in enumerate(failed_scenarios, 1):
        expected_text = failed_scenario['expected_text']
        total_audio = failed_scenario['total_audio']
        audio_details = failed_scenario['audio_details']
        
        merge_start_row = row_num
        merge_end_row = row_num + len(audio_details) - 1
        
        detail_ws.cell(row=row_num, column=1, value=idx)
        detail_ws.cell(row=row_num, column=2, value=expected_text)
        detail_ws.cell(row=row_num, column=3, value=total_audio)
        detail_ws.cell(row=row_num, column=4, value=audio_details[0]['voice'] if audio_details else '')
        detail_ws.cell(row=row_num, column=5, value=audio_details[0]['actual'] if audio_details else '')
        detail_ws.cell(row=row_num, column=6, value=audio_details[0]['result'] if audio_details else '')
        detail_ws.cell(row=row_num, column=7, value=audio_details[0]['confidence'] if audio_details else '')
        
        for col_idx in range(1, len(detail_headers) + 1):
            cell = detail_ws.cell(row=row_num, column=col_idx)
            cell.border = border
            cell.alignment = alignment
        
        row_num += 1
        
        for audio_idx in range(1, len(audio_details)):
            detail_ws.cell(row=row_num, column=1, value='')
            detail_ws.cell(row=row_num, column=2, value='')
            detail_ws.cell(row=row_num, column=3, value='')
            detail_ws.cell(row=row_num, column=4, value=audio_details[audio_idx]['voice'])
            detail_ws.cell(row=row_num, column=5, value=audio_details[audio_idx]['actual'])
            detail_ws.cell(row=row_num, column=6, value=audio_details[audio_idx]['result'])
            detail_ws.cell(row=row_num, column=7, value=audio_details[audio_idx]['confidence'])
            
            for col_idx in range(1, len(detail_headers) + 1):
                cell = detail_ws.cell(row=row_num, column=col_idx)
                cell.border = border
                cell.alignment = alignment
            
            row_num += 1
        
        if merge_end_row > merge_start_row:
            detail_ws.merge_cells(f'A{merge_start_row}:A{merge_end_row}')
            detail_ws.merge_cells(f'B{merge_start_row}:B{merge_end_row}')
            detail_ws.merge_cells(f'C{merge_start_row}:C{merge_end_row}')
        
        row_num += 1


def setup_header_style(ws, row_num, columns):
    """
    设置表头样式
    
    Args:
        ws: worksheet对象
        row_num: 行号
        columns: 列宽字典
    """
    # 定义样式
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # 应用样式
    for col_letter, width in columns.items():
        col_idx = list(columns.keys()).index(col_letter) + 1
        cell = ws.cell(row=row_num, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = alignment
        ws.column_dimensions[col_letter].width = width


def generate_report(all_results, output_file='report.xlsx'):
    """
    生成Excel报告文件
    
    Args:
        all_results: 所有文件的分析结果，格式为 {file_path: (scenario_stats, failed_scenarios)}
        output_file: 输出文件路径
    """
    wb = Workbook()
    
    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # 创建汇总信息表
    summary_ws = wb.create_sheet("汇总信息", 0)
    
    # 设置汇总表表头
    summary_headers = ['场景名称', '文件路径', '总场景数', '识别正确', '识别错误', '通过率(%)', '全错场景数']
    summary_columns = {'A': 20, 'B': 60, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 14}
    
    for idx, header in enumerate(summary_headers, 1):
        summary_ws.cell(row=1, column=idx, value=header)
    
    setup_header_style(summary_ws, 1, summary_columns)
    
    # 总体统计
    total_scenarios = 0
    total_correct = 0
    total_failed = 0
    total_corrected = 0  # 通过大小写兼容纠正的总数
    
    # 遍历每个文件，填充汇总表
    row_num = 2
    
    for file_path, (scenario_stats, failed_scenarios) in all_results.items():
        scenario_name = get_scenario_name(file_path)
        
        # 统计该场景的数据
        file_total = len(scenario_stats)
        file_correct = sum(1 for stats in scenario_stats.values() if stats['is_correct'])
        file_failed = file_total - file_correct
        file_pass_rate = (file_correct / file_total * 100) if file_total > 0 else 0
        file_corrected = sum(stats.get('corrected', 0) for stats in scenario_stats.values())
        
        total_scenarios += file_total
        total_correct += file_correct
        total_failed += file_failed
        total_corrected += file_corrected
        
        # 填充汇总表数据
        summary_ws.cell(row=row_num, column=1, value=scenario_name)
        summary_ws.cell(row=row_num, column=2, value=file_path)
        summary_ws.cell(row=row_num, column=3, value=file_total)
        summary_ws.cell(row=row_num, column=4, value=file_correct)
        summary_ws.cell(row=row_num, column=5, value=file_failed)
        summary_ws.cell(row=row_num, column=6, value=round(file_pass_rate, 2))
        summary_ws.cell(row=row_num, column=7, value=len(failed_scenarios))
        
        # 设置数据行样式
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        for col_idx in range(1, len(summary_headers) + 1):
            cell = summary_ws.cell(row=row_num, column=col_idx)
            cell.border = border
            cell.alignment = alignment
        
        row_num += 1
        
        # 如果有全错场景，创建该场景的详细错误报告sheet并立即填充数据
        if failed_scenarios:
            # Excel sheet名称不能超过31个字符，且不能包含某些特殊字符
            sheet_name = scenario_name[:31].replace('/', '_').replace('\\', '_').replace('?', '_').replace('*', '_').replace('[', '_').replace(']', '_')
            # 如果sheet名称已存在，添加序号
            original_sheet_name = sheet_name
            counter = 1
            while sheet_name in wb.sheetnames:
                sheet_name = f"{original_sheet_name[:28]}_{counter}"
                counter += 1
            
            detail_ws = wb.create_sheet(sheet_name)
            fill_detail_sheet(detail_ws, failed_scenarios)
    
    # 添加总体统计行
    row_num += 1
    summary_ws.cell(row=row_num, column=1, value="总计")
    summary_ws.cell(row=row_num, column=2, value="")
    summary_ws.cell(row=row_num, column=3, value=total_scenarios)
    summary_ws.cell(row=row_num, column=4, value=total_correct)
    summary_ws.cell(row=row_num, column=5, value=total_failed)
    if total_scenarios > 0:
        overall_pass_rate = (total_correct / total_scenarios * 100)
        summary_ws.cell(row=row_num, column=6, value=round(overall_pass_rate, 2))
    else:
        summary_ws.cell(row=row_num, column=6, value=0)
    summary_ws.cell(row=row_num, column=7, value="")
    
    # 设置总计行样式（加粗）
    bold_font = Font(bold=True)
    for col_idx in range(1, len(summary_headers) + 1):
        cell = summary_ws.cell(row=row_num, column=col_idx)
        cell.font = bold_font
    
    # 添加说明行
    row_num += 2
    summary_ws.cell(row=row_num, column=1, value="说明：")
    summary_ws.merge_cells(f'A{row_num}:G{row_num}')
    summary_ws.cell(row=row_num, column=1, value=f"说明：每个场景对应多个音频，只要有一个音频识别成功，该场景就算正确。通过大小写兼容纠正了 {total_corrected} 个识别结果。")
    summary_ws.cell(row=row_num, column=1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # 保存文件
    wb.save(output_file)


def main():
    """主函数"""
    # 创建命令行参数解析器
    parser = argparse.ArgumentParser(
        description='测试结果分析工具 - 用于分析result.csv或result.txt文件，统计识别率并生成报告',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
使用示例:
  python asr_analyze_results.py /path/to/search /path/to/output/report.xlsx
  python asr_analyze_results.py -s /path/to/search -o /path/to/output/report.xlsx
  
说明:
  - 优先查找result.csv文件
  - 如果找不到result.csv，则查找result.txt文件（需要验证第一行格式）
        '''
    )
    
    parser.add_argument(
        'search_dir',
        nargs='?',
        help='查找result.csv或result.txt文件的目录路径'
    )
    
    parser.add_argument(
        'output_file',
        nargs='?',
        help='报告输出文件路径（Excel格式）'
    )
    
    parser.add_argument(
        '-s', '--search-dir',
        dest='search_dir_alt',
        help='查找result.csv或result.txt文件的目录路径（可选，与位置参数二选一）'
    )
    
    parser.add_argument(
        '-o', '--output',
        dest='output_file_alt',
        help='报告输出文件路径（可选，与位置参数二选一）'
    )
    
    args = parser.parse_args()
    
    # 确定查找目录（优先使用 -s 参数，其次使用位置参数）
    search_dir = args.search_dir_alt or args.search_dir
    output_file = args.output_file_alt or args.output_file
    
    # 验证输入路径
    if not search_dir:
        parser.error("必须指定查找目录路径（使用位置参数或 -s/--search-dir 选项）")
    
    if not output_file:
        parser.error("必须指定输出文件路径（使用位置参数或 -o/--output 选项）")
    
    # 检查查找目录是否存在
    if not os.path.exists(search_dir):
        parser.error(f"查找目录不存在: {search_dir}")
    
    if not os.path.isdir(search_dir):
        parser.error(f"输入路径不是目录: {search_dir}")
    
    # 确保输出目录存在
    output_dir = os.path.dirname(output_file)
    if output_dir and not os.path.exists(output_dir):
        try:
            os.makedirs(output_dir, exist_ok=True)
            print(f"已创建输出目录: {output_dir}")
        except Exception as e:
            parser.error(f"无法创建输出目录 {output_dir}: {e}")
    
    print(f"开始查找result.csv或result.txt文件，查找目录: {search_dir}")
    
    # 查找所有result文件（优先csv，找不到则找txt）
    result_files = find_all_result_files(search_dir)
    
    if not result_files:
        print("未找到任何result.csv或result.txt文件！")
        return
    
    print(f"找到 {len(result_files)} 个结果文件:")
    for file_path, file_type in result_files:
        print(f"  - [{file_type.upper()}] {file_path}")
    
    # 分析每个文件
    all_results = {}
    
    for file_path, file_type in result_files:
        print(f"\n正在分析: [{file_type.upper()}] {file_path}")
        data = read_result_file(file_path, file_type)
        
        if not data:
            print(f"  警告: {file_path} 文件为空或无法读取")
            continue
        
        scenario_stats, failed_scenarios = analyze_scenario_results(data)
        all_results[file_path] = (scenario_stats, failed_scenarios)
        
        print(f"  场景数: {len(scenario_stats)}, 全错场景数: {len(failed_scenarios)}")
    
    # 生成报告
    print(f"\n正在生成报告: {output_file}")
    generate_report(all_results, output_file)
    
    print(f"\n报告生成完成！报告文件: {output_file}")


if __name__ == '__main__':
    main()

