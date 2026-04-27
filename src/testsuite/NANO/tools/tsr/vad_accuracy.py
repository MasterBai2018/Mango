#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# @Time    : 2026/03/04
# @Author  : huidong.bai
# @File    : vad_accuracy.py
# @Software: PyCharm
# @Mail    : MasterBai2018@outlook.com
"""
VAD_ACCURACY 指令实现

用于计算 VAD（语音活动检测）的 FA/FR 准确率，支持：
- False Acceptance (FA) 误检率
- False Rejection  (FR) 漏检率
- 断句 / 连句统计
- 每条音频检测率

判定规则：
1. 时间重叠 = 正确匹配（VAD 无关键词概念，仅判断时间是否重叠）
2. 无重叠：ref 有但 result 没有 = FR（漏检）；result 有但 ref 没有 = FA（误检）
3. 断句：1 个 ref 段对应 ≥2 个 result 段且全部重叠（不计 FA/FR）
4. 连句：≥2 个 ref 段对应 1 个 result 段且全部重叠（不计 FA/FR）

输入文件格式（ref 和 result 相同格式，空格分隔，无表头）：
    音频路径  00h_S  开始时间(s)  结束时间(s)

使用示例:
    [TSR]VAD_ACCURACY result=vad.txt ref=ref.txt [output=report.xlsx]
"""
import os
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Tuple, Optional
from loguru import logger

try:
    from . import register_tsr_command
    from .base import BaseTSRHandler
except ImportError:
    _current_dir = os.path.dirname(os.path.abspath(__file__))
    _project_root = os.path.abspath(os.path.join(_current_dir, '..', '..', '..', '..', '..'))
    if _project_root not in sys.path:
        sys.path.insert(0, _project_root)
    from src.testsuite.NANO.tools.tsr import register_tsr_command
    from src.testsuite.NANO.tools.tsr.base import BaseTSRHandler


# ─── 数据结构 ──────────────────────────────────────────────────────────────────

class VADSegment:
    """VAD 时间段"""

    def __init__(self, audio_path: str, start: float, end: float, line_num: int = 0):
        self.audio_path = audio_path
        self.start      = start
        self.end        = end
        self.line_num   = line_num
        self.matched    = False

    def has_overlap(self, other: 'VADSegment') -> bool:
        """判断两个时间段是否有重叠"""
        return not (other.start >= self.end or other.end <= self.start)

    def __repr__(self) -> str:
        return f"VADSegment({self.audio_path}, {self.start:.3f}, {self.end:.3f})"


@dataclass
class VADAccuracyStats:
    """VAD 准确率汇总统计"""
    total_ref:      int = 0
    total_result:   int = 0
    matched_count:  int = 0
    fa_count:       int = 0
    fr_count:       int = 0
    break_count:    int = 0
    join_count:     int = 0
    is_fa_only:     bool = False   # ref=NULL 时为 True，表示 FA-only 场景

    fa_segments:    List[VADSegment]  = field(default_factory=list)
    fr_segments:    List[VADSegment]  = field(default_factory=list)
    break_cases:    List[Dict]        = field(default_factory=list)
    join_cases:     List[Dict]        = field(default_factory=list)
    matched_pairs:  List[Tuple]       = field(default_factory=list)

    all_ref_segments:    List[VADSegment] = field(default_factory=list)
    all_result_segments: List[VADSegment] = field(default_factory=list)

    audio_stats:    Dict[str, Dict]   = field(default_factory=dict)

    @property
    def fa_rate(self) -> float:
        """误检率"""
        return self.fa_count / self.total_result if self.total_result > 0 else 0.0

    @property
    def fr_rate(self) -> float:
        """漏检率"""
        return self.fr_count / self.total_ref if self.total_ref > 0 else 0.0

    @property
    def detection_rate(self) -> float:
        """检测率 = (ref 总数 - FR 数) / ref 总数"""
        return (self.total_ref - self.fr_count) / self.total_ref if self.total_ref > 0 else 0.0


# ─── 文件解析 ──────────────────────────────────────────────────────────────────

def _parse_vad_file(file_path: str) -> List[VADSegment]:
    """
    解析 VAD 文件（ref 和 result 使用相同格式）
    格式：音频路径  00h_S  开始时间(s)  结束时间(s)（空格分隔，无表头）
    """
    segments: List[VADSegment] = []
    with open(file_path, 'r', encoding='utf-8') as f:
        for line_num, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            parts = line.split()
            if len(parts) < 4:
                logger.warning(f"VAD 文件第 {line_num} 行格式不正确，跳过：{line}")
                continue
            try:
                audio_path = parts[0]
                # parts[1] 为固定标记（如 "00h_S"），跳过
                start_time = float(parts[2])
                end_time   = float(parts[3])
                if start_time >= end_time:
                    logger.warning(f"VAD 文件第 {line_num} 行 start >= end，跳过")
                    continue
                segments.append(VADSegment(audio_path, start_time, end_time, line_num))
            except (ValueError, IndexError) as e:
                logger.warning(f"VAD 文件第 {line_num} 行解析失败，跳过：{e}")
    return segments


def _group_by_audio(segments: List[VADSegment]) -> Dict[str, List[VADSegment]]:
    """按音频路径分组并按开始时间排序"""
    grouped: Dict[str, List[VADSegment]] = defaultdict(list)
    for seg in segments:
        grouped[seg.audio_path].append(seg)
    for segs in grouped.values():
        segs.sort(key=lambda s: (s.start, s.end))
    return grouped


# ─── 核心匹配算法（修复版：支持多段音频内逐段断句/连句检测）──────────────────────

def _process_audio_segments(
    ref_segs:    List[VADSegment],
    result_segs: List[VADSegment],
) -> Tuple[
    List[Tuple[VADSegment, VADSegment]],  # matched_pairs
    List[VADSegment],                      # fa_segments
    List[VADSegment],                      # fr_segments
    List[Dict],                            # break_cases
    List[Dict],                            # join_cases
]:
    """
    处理单个音频的 VAD 段匹配（修复版）

    匹配优先级：
      1. 断句检测（per-segment：1 ref → ≥2 result，全部时间重叠）
      2. 连句检测（per-segment：≥2 ref → 1 result，全部时间重叠）
      3. 常规 1-to-1 匹配（时间有重叠）
      4. 剩余 ref → FR；剩余 result → FA
    """
    for seg in ref_segs:    seg.matched = False
    for seg in result_segs: seg.matched = False

    matched_pairs: List[Tuple] = []
    fa_segments:   List[VADSegment] = []
    fr_segments:   List[VADSegment] = []
    break_cases:   List[Dict] = []
    join_cases:    List[Dict] = []

    # ── Step 1：逐段断句检测 ──────────────────────────────────────────────────
    for ref_seg in ref_segs:
        if ref_seg.matched:
            continue
        overlap_indices = [
            i for i, r in enumerate(result_segs)
            if not r.matched and ref_seg.has_overlap(r)
        ]
        if len(overlap_indices) >= 2:
            for idx in overlap_indices:
                result_segs[idx].matched = True
            ref_seg.matched = True
            break_cases.append({
                'audio_path':      ref_seg.audio_path,
                'ref':             ref_seg,
                'result_segments': [result_segs[i] for i in overlap_indices],
                'result_indices':  overlap_indices,
            })

    # ── Step 2：逐段连句检测 ──────────────────────────────────────────────────
    for result_seg in result_segs:
        if result_seg.matched:
            continue
        overlap_indices = [
            i for i, r in enumerate(ref_segs)
            if not r.matched and result_seg.has_overlap(r)
        ]
        if len(overlap_indices) >= 2:
            for idx in overlap_indices:
                ref_segs[idx].matched = True
            result_seg.matched = True
            join_cases.append({
                'audio_path':   result_seg.audio_path,
                'ref_segments': [ref_segs[i] for i in overlap_indices],
                'ref_indices':  overlap_indices,
                'result':       result_seg,
            })

    # ── Step 3：常规 1-to-1 匹配 ─────────────────────────────────────────────
    for ref_seg in ref_segs:
        if ref_seg.matched:
            continue
        match_idx = next(
            (i for i, r in enumerate(result_segs)
             if not r.matched and ref_seg.has_overlap(r)),
            None
        )
        if match_idx is not None:
            result_segs[match_idx].matched = True
            ref_seg.matched = True
            matched_pairs.append((ref_seg, result_segs[match_idx]))
        else:
            fr_segments.append(ref_seg)

    # ── Step 4：剩余未匹配 result → FA ───────────────────────────────────────
    for result_seg in result_segs:
        if not result_seg.matched:
            fa_segments.append(result_seg)

    return matched_pairs, fa_segments, fr_segments, break_cases, join_cases


def _calculate_stats(ref_file: str, result_file: str) -> VADAccuracyStats:
    """读取文件、分组、逐音频匹配，汇总统计结果"""
    # ref=NULL 表示 FA-only 场景，不解析参考文件
    is_fa_only = ref_file.upper() == 'NULL'
    if is_fa_only:
        ref_segments = []
        logger.info("FA-only 模式：ref=NULL，仅统计误检（FA），跳过参考文件解析")
    else:
        ref_segments = _parse_vad_file(ref_file)
        logger.info(f"ref 文件共 {len(ref_segments)} 条记录")

    result_segments = _parse_vad_file(result_file)
    logger.info(f"result 文件共 {len(result_segments)} 条记录")

    ref_grouped    = _group_by_audio(ref_segments)
    result_grouped = _group_by_audio(result_segments)
    all_audio_paths = set(ref_grouped.keys()) | set(result_grouped.keys())

    stats = VADAccuracyStats(
        total_ref=len(ref_segments),
        total_result=len(result_segments),
        all_ref_segments=ref_segments,
        all_result_segments=result_segments,
        is_fa_only=is_fa_only,
    )

    for audio_path in sorted(all_audio_paths):
        ref_segs    = ref_grouped.get(audio_path, [])
        result_segs = result_grouped.get(audio_path, [])

        mp, fa, fr, bk, jn = _process_audio_segments(ref_segs, result_segs)

        stats.matched_pairs.extend(mp)
        stats.fa_segments.extend(fa)
        stats.fr_segments.extend(fr)
        stats.break_cases.extend(bk)
        stats.join_cases.extend(jn)

        expected    = len(ref_segs)
        fr_n        = len(fr)
        fa_n        = len(fa)
        det_rate    = ((expected - fr_n) / expected * 100.0) if expected > 0 else 0.0
        stats.audio_stats[audio_path] = {
            'expected':       expected,
            'actual':         len(result_segs),
            'fr':             fr_n,
            'fa':             fa_n,
            'detection_rate': det_rate,
        }

    stats.matched_count = len(stats.matched_pairs)
    stats.fa_count      = len(stats.fa_segments)
    stats.fr_count      = len(stats.fr_segments)
    stats.break_count   = len(stats.break_cases)
    stats.join_count    = len(stats.join_cases)

    return stats


# ─── TSR 指令处理器 ───────────────────────────────────────────────────────────

@register_tsr_command("VAD_ACCURACY")
class VADAccuracyHandler(BaseTSRHandler):
    """
    VAD FA/FR 准确率计算处理器

    指令格式:
        [TSR]VAD_ACCURACY result=<result_txt> ref=<ref_txt> [output=<report_xlsx>]

    参数说明:
        result : VAD 识别结果文件路径（必需）
        ref    : 参考 VAD 文件路径（必需）
        output : Excel 报告输出路径（可选）

    输入文件格式（ref/result 相同）：
        音频路径  00h_S  开始时间(s)  结束时间(s)
    """

    COMMAND_NAME = "VAD_ACCURACY"

    def execute(self, params: list, output_report: str = None) -> str:
        """执行 VAD FA/FR 计算"""
        result_file, ref_file, xlsx_path, output_was_specified = self._parse_params(
            params, output_report
        )

        if not os.path.exists(result_file):
            raise FileNotFoundError(f"VAD 结果文件不存在: {result_file}")
        # ref="NULL" 为 FA-only 模式（无参考文件，仅统计误检）；否则文件必须存在
        if ref_file.upper() != 'NULL' and not os.path.exists(ref_file):
            raise FileNotFoundError(f"参考 VAD 文件不存在: {ref_file}")

        logger.info(f"VAD_ACCURACY: result={result_file}, ref={ref_file}")

        stats = _calculate_stats(ref_file, result_file)

        # 生成 Excel 报告
        self._generate_excel_report(stats, ref_file, result_file, xlsx_path)

        # 生成 HTML 报告（保存到 allure_result/mango_report/{define}/{suite}/）
        html_content  = self._build_vad_html(stats)
        html_filename = "vad_accuracy.html"
        self.save_html_report(html_content, html_filename, "VAD FA/FR 报告", "vad_accuracy")

        # 本地副本
        local_html = os.path.splitext(xlsx_path)[0] + '.html'
        os.makedirs(os.path.dirname(local_html) or '.', exist_ok=True)
        with open(local_html, 'w', encoding='utf-8') as f:
            f.write(html_content)
        logger.info(f"本地 HTML 报告已生成: {local_html}")

        return self._generate_summary(stats, xlsx_path)

    # ── 参数解析 ──────────────────────────────────────────────────────────────

    def _parse_params(
        self, params: list, default_output: str = None
    ) -> Tuple[str, str, str, bool]:
        result_file = None
        ref_file    = None
        output_path = default_output

        for param in params:
            if   param.startswith("result="):
                result_file = param[7:]
            elif param.startswith("ref="):
                ref_file    = param[4:]
            elif param.startswith("output="):
                output_path = param[7:]

        if not result_file:
            raise ValueError("缺少必需参数: result=<result_txt>")
        if not ref_file:
            raise ValueError("缺少必需参数: ref=<ref_txt>")

        output_was_specified = output_path is not None
        if not output_path:
            suite_mango_dir = self._get_suite_mango_dir()
            output_path     = os.path.join(suite_mango_dir, "vad_accuracy.xlsx")

        return result_file, ref_file, output_path, output_was_specified

    # ── HTML 报告 ─────────────────────────────────────────────────────────────

    def generate_html_report(self, *args, **kwargs) -> Optional[str]:
        """实现抽象方法"""
        stats = args[0] if args else kwargs.get('stats')
        if stats is None:
            return None
        html = self._build_vad_html(stats)
        return self.save_html_report(html, "vad_accuracy.html", "VAD FA/FR 报告", "vad_accuracy")

    def _build_vad_html(self, stats: VADAccuracyStats) -> str:
        """构建自包含 HTML"""
        fa_cls  = ('card-success' if stats.fa_rate <= 0.05 else
                   'card-warning' if stats.fa_rate <= 0.10 else 'card-failure')

        if stats.is_fa_only:
            # FA-only 模式：仅展示 Result 总数和 FA 误检数
            cards_html = (
                '<div class="summary-cards">'
                '<div class="card card-warning"><div class="card-value">FA-Only</div>'
                '<div class="card-label">运行模式</div></div>'
                f'<div class="card"><div class="card-value">{stats.total_result}</div>'
                f'<div class="card-label">Result 总数</div></div>'
                f'<div class="card {fa_cls}"><div class="card-value">{stats.fa_count}</div>'
                f'<div class="card-label">FA 误检数</div></div>'
                f'</div>'
            )
        else:
            det_cls = ('card-success' if stats.detection_rate >= 0.95 else
                       'card-warning' if stats.detection_rate >= 0.80 else 'card-failure')
            fr_cls  = ('card-success' if stats.fr_rate <= 0.05 else
                       'card-warning' if stats.fr_rate <= 0.10 else 'card-failure')
            cards_html = (
                '<div class="summary-cards">'
                f'<div class="card"><div class="card-value">{stats.total_ref}</div>'
                f'<div class="card-label">Ref 总数</div></div>'
                f'<div class="card"><div class="card-value">{stats.total_result}</div>'
                f'<div class="card-label">Result 总数</div></div>'
                f'<div class="card card-success"><div class="card-value">{stats.matched_count}</div>'
                f'<div class="card-label">匹配数</div></div>'
                f'<div class="card {det_cls}"><div class="card-value">{stats.detection_rate * 100:.2f}%</div>'
                f'<div class="card-label">检测率</div></div>'
                f'<div class="card {fa_cls}"><div class="card-value">{stats.fa_rate * 100:.2f}%</div>'
                f'<div class="card-label">FA 误检率</div></div>'
                f'<div class="card {fr_cls}"><div class="card-value">{stats.fr_rate * 100:.2f}%</div>'
                f'<div class="card-label">FR 漏检率</div></div>'
                f'<div class="card card-warning"><div class="card-value">{stats.break_count}</div>'
                f'<div class="card-label">断句数</div></div>'
                f'<div class="card card-warning"><div class="card-value">{stats.join_count}</div>'
                f'<div class="card-label">连句数</div></div>'
                f'</div>'
            )

        # Tab 1：每条音频统计
        audio_rows = ''
        for ap in sorted(stats.audio_stats.keys()):
            s   = stats.audio_stats[ap]
            dr  = s['detection_rate']
            cls = ('tag-fail' if dr < 80 else 'tag-pass' if dr >= 95 else 'tag-warn')
            audio_rows += (
                f'<tr><td>{os.path.basename(ap)}</td>'
                f'<td>{s["expected"]}</td><td>{s["actual"]}</td>'
                f'<td>{s["fr"]}</td><td>{s["fa"]}</td>'
                f'<td><span class="{cls}">{dr:.2f}%</span></td></tr>'
            )
        tab1_html = (
            '<table class="report-table"><thead>'
            '<tr><th>音频文件</th><th>预期检测数</th><th>实际检测数</th>'
            '<th>FR</th><th>FA</th><th>检测率</th></tr>'
            f'</thead><tbody>{audio_rows}</tbody></table>'
        )

        # Tab 2：FA 详情
        fa_rows = ''.join(
            f'<tr><td>{i}</td><td>{os.path.basename(s.audio_path)}</td>'
            f'<td>{s.start:.3f}</td><td>{s.end:.3f}</td>'
            f'<td>{s.end - s.start:.3f}</td></tr>'
            for i, s in enumerate(stats.fa_segments, 1)
        )
        tab2_html = (
            '<table class="report-table"><thead>'
            '<tr><th>#</th><th>音频</th><th>开始(s)</th><th>结束(s)</th><th>时长(s)</th></tr>'
            f'</thead><tbody>{fa_rows}</tbody></table>'
        )

        # Tab 3：FR 详情
        fr_rows = ''.join(
            f'<tr><td>{i}</td><td>{os.path.basename(s.audio_path)}</td>'
            f'<td>{s.start:.3f}</td><td>{s.end:.3f}</td>'
            f'<td>{s.end - s.start:.3f}</td></tr>'
            for i, s in enumerate(stats.fr_segments, 1)
        )
        tab3_html = (
            '<table class="report-table"><thead>'
            '<tr><th>#</th><th>音频</th><th>开始(s)</th><th>结束(s)</th><th>时长(s)</th></tr>'
            f'</thead><tbody>{fr_rows}</tbody></table>'
        )

        section_html = (
            '<div class="section">'
            '<div class="section-header">📋 VAD 检测详情</div>'
            '<div class="section-body">'
            '<div class="inner-tabs">'
            '<div class="inner-tab active" onclick="switchTab(this,\'va1\')">音频统计</div>'
            f'<div class="inner-tab" onclick="switchTab(this,\'va2\')">FA 误检 ({stats.fa_count})</div>'
            f'<div class="inner-tab" onclick="switchTab(this,\'va3\')">FR 漏检 ({stats.fr_count})</div>'
            '</div>'
            f'<div id="va1" class="inner-pane active">{tab1_html}</div>'
            f'<div id="va2" class="inner-pane">{tab2_html}</div>'
            f'<div id="va3" class="inner-pane">{tab3_html}</div>'
            '</div></div>'
        )

        return self.build_html_page("VAD FA/FR 报告", cards_html + section_html)

    # ── Excel 报告（5 个 Sheet） ───────────────────────────────────────────────

    def _generate_excel_report(
        self,
        stats:       VADAccuracyStats,
        ref_file:    str,
        result_file: str,
        output_path: str,
    ) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("生成 Excel 报告需要安装 openpyxl：pip install openpyxl")

        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

        header_fill  = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font  = Font(bold=True, color="FFFFFF", size=11)
        title_font   = Font(bold=True, size=12)
        center_align = Alignment(horizontal='center', vertical='center')
        left_align   = Alignment(horizontal='left',   vertical='center')
        red_fill     = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        green_fill   = PatternFill(start_color="C5E0B4", end_color="C5E0B4", fill_type="solid")
        blue_fill    = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

        def _write_header(ws, row: int, headers: list) -> None:
            for c, h in enumerate(headers, 1):
                cell           = ws.cell(row=row, column=c, value=h)
                cell.fill      = header_fill
                cell.font      = header_font
                cell.alignment = center_align

        # ── Sheet 1：概要 ──────────────────────────────────────────────────────
        ws_sum = wb.create_sheet("概要", 0)
        ws_sum['A1'] = 'VAD FA/FR 分析结果'
        ws_sum['A1'].font = Font(bold=True, size=14)
        ws_sum.merge_cells('A1:B1')

        r = 3
        for label, value in [('参考文件：', ref_file), ('结果文件：', result_file)]:
            ws_sum.cell(row=r, column=1, value=label)
            ws_sum.cell(row=r, column=2, value=value)
            r += 1

        r += 1
        ws_sum.cell(row=r, column=1, value='汇总统计').font = title_font
        ws_sum.merge_cells(f'A{r}:B{r}')
        r += 1

        total_det_rate = stats.detection_rate * 100.0
        for label, value in [
            ('Ref 总段数',             stats.total_ref),
            ('Result 总段数',          stats.total_result),
            ('成功匹配段数',           stats.matched_count),
            ('断句数量',               stats.break_count),
            ('连句数量',               stats.join_count),
            ('', ''),
            ('False Acceptance (FA)',   f'{stats.fa_count} 段 ({stats.fa_rate * 100:.2f}%)'),
            ('False Rejection (FR)',    f'{stats.fr_count} 段 ({stats.fr_rate * 100:.2f}%)'),
        ]:
            cell_a = ws_sum.cell(row=r, column=1, value=label)
            cell_b = ws_sum.cell(row=r, column=2, value=value)
            if label in ('False Acceptance (FA)', 'False Rejection (FR)'):
                cell_a.font = Font(bold=True)
                cell_b.font = Font(bold=True)
            r += 1

        ws_sum.cell(row=r, column=1, value='总检测率')
        ws_sum.cell(row=r, column=2, value=f'{total_det_rate:.2f}%')
        ws_sum.cell(row=r, column=1).font = Font(bold=True, color="0000FF")
        ws_sum.cell(row=r, column=2).font = Font(bold=True, color="0000FF")
        r += 2

        # 每条音频统计
        ws_sum.cell(row=r, column=1, value='每条音频检测统计详情').font = title_font
        ws_sum.merge_cells(f'A{r}:F{r}')
        r += 1
        _write_header(ws_sum, r, ['音频路径', '预期检测数', '实际检测数', 'FR 数量', 'FA 数量', '检测率'])
        r += 1

        for ap in sorted(stats.audio_stats.keys()):
            s   = stats.audio_stats[ap]
            dr  = s['detection_rate']
            row_fill = (red_fill if dr < 80 else green_fill if dr >= 95 else None)
            ws_sum.cell(row=r, column=1, value=ap).alignment = left_align
            ws_sum.cell(row=r, column=2, value=s['expected']).alignment  = center_align
            ws_sum.cell(row=r, column=3, value=s['actual']).alignment    = center_align
            ws_sum.cell(row=r, column=4, value=s['fr']).alignment        = center_align
            ws_sum.cell(row=r, column=5, value=s['fa']).alignment        = center_align
            ws_sum.cell(row=r, column=6, value=f"{dr:.2f}%").alignment   = center_align
            if row_fill:
                for col in range(1, 7):
                    ws_sum.cell(row=r, column=col).fill = row_fill
            r += 1

        # 总计行
        tot_exp    = sum(s['expected'] for s in stats.audio_stats.values())
        tot_actual = sum(s['actual']   for s in stats.audio_stats.values())
        tot_fr     = sum(s['fr']       for s in stats.audio_stats.values())
        tot_fa     = sum(s['fa']       for s in stats.audio_stats.values())
        for c, (v, align) in enumerate([
            ('总计', center_align), (tot_exp, center_align), (tot_actual, center_align),
            (tot_fr, center_align), (tot_fa, center_align),
            (f'{total_det_rate:.2f}%', center_align)
        ], 1):
            cell           = ws_sum.cell(row=r, column=c, value=v)
            cell.font      = Font(bold=True, color="0000FF") if c == 6 else Font(bold=True)
            cell.fill      = blue_fill
            cell.alignment = align

        ws_sum.column_dimensions['A'].width = 60
        for ltr, w in zip('BCDEF', [15, 15, 12, 12, 12]):
            ws_sum.column_dimensions[ltr].width = w

        # ── Sheet 2 / 3：FA / FR 详情 ──────────────────────────────────────────
        seg_hdrs = ['音频路径', '开始时间', '结束时间', '时长(秒)', '行号']
        for sheet_idx, (title, segments) in enumerate([
            ("FA详情", stats.fa_segments),
            ("FR详情", stats.fr_segments),
        ], 1):
            ws = wb.create_sheet(title, sheet_idx)
            _write_header(ws, 1, seg_hdrs)
            for ri, seg in enumerate(segments, 2):
                ws.cell(row=ri, column=1, value=seg.audio_path)
                ws.cell(row=ri, column=2, value=seg.start)
                ws.cell(row=ri, column=3, value=seg.end)
                ws.cell(row=ri, column=4, value=round(seg.end - seg.start, 6))
                ws.cell(row=ri, column=5, value=seg.line_num)
            for ci, w in enumerate([60, 15, 15, 12, 10], 1):
                ws.column_dimensions[get_column_letter(ci)].width = w
            ws.freeze_panes = 'A2'

        # ── Sheet 4：Signal 全览 ───────────────────────────────────────────────
        ws_sig = wb.create_sheet("Signal", 3)
        _write_header(ws_sig, 1, ['音频路径', '开始时间', '结束时间', '时长(秒)', '行号', '状态'])

        matched_result_set = {res for _, res in stats.matched_pairs}
        fa_result_set      = set(stats.fa_segments)
        break_result_set   = {seg for case in stats.break_cases for seg in case['result_segments']}
        join_result_set    = {case['result'] for case in stats.join_cases}

        for ri, seg in enumerate(stats.all_result_segments, 2):
            if   seg in matched_result_set: status, fill = '匹配',  green_fill
            elif seg in fa_result_set:      status, fill = 'FA',    red_fill
            elif seg in break_result_set:   status, fill = '断句',  None
            elif seg in join_result_set:    status, fill = '连句',  None
            else:                           status, fill = '其他',  None
            for ci, v in enumerate([seg.audio_path, seg.start, seg.end,
                                    round(seg.end - seg.start, 6), seg.line_num, status], 1):
                cell = ws_sig.cell(row=ri, column=ci, value=v)
                if fill:
                    cell.fill = fill
        for ci, w in enumerate([60, 15, 15, 12, 10, 10], 1):
            ws_sig.column_dimensions[get_column_letter(ci)].width = w
        ws_sig.freeze_panes = 'A2'

        # ── Sheet 5：Ref 全览 ──────────────────────────────────────────────────
        ws_ref_sheet = wb.create_sheet("Ref", 4)
        _write_header(ws_ref_sheet, 1, ['音频路径', '开始时间', '结束时间', '时长(秒)', '行号', '状态'])

        matched_ref_set = {ref for ref, _ in stats.matched_pairs}
        fr_ref_set      = set(stats.fr_segments)
        break_ref_set   = {case['ref'] for case in stats.break_cases}
        join_ref_set    = {seg for case in stats.join_cases for seg in case['ref_segments']}

        for ri, seg in enumerate(stats.all_ref_segments, 2):
            if   seg in matched_ref_set: status, fill = '匹配',  green_fill
            elif seg in fr_ref_set:      status, fill = 'FR',    red_fill
            elif seg in break_ref_set:   status, fill = '断句',  None
            elif seg in join_ref_set:    status, fill = '连句',  None
            else:                        status, fill = '其他',  None
            for ci, v in enumerate([seg.audio_path, seg.start, seg.end,
                                    round(seg.end - seg.start, 6), seg.line_num, status], 1):
                cell = ws_ref_sheet.cell(row=ri, column=ci, value=v)
                if fill:
                    cell.fill = fill
        for ci, w in enumerate([60, 15, 15, 12, 10, 10], 1):
            ws_ref_sheet.column_dimensions[get_column_letter(ci)].width = w
        ws_ref_sheet.freeze_panes = 'A2'

        parent = os.path.dirname(output_path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        wb.save(output_path)
        logger.info(f"Excel 报告已生成: {output_path}")

    # ── Summary 文本 ──────────────────────────────────────────────────────────

    def _generate_summary(self, stats: VADAccuracyStats, output_path: str) -> str:
        if stats.is_fa_only:
            # FA-only 模式：无参考文件，仅展示误检统计
            content_lines = [
                f"*  [FA-only 模式] 无参考文件，仅统计误检（FA）",
                f"",
                f"*  Result 总数:   {stats.total_result}",
                f"*  FA 误检数:     {stats.fa_count}",
                f"",
                f"*  详细报告: {output_path}",
            ]
        else:
            content_lines = [
                f"*  Ref 总数:   {stats.total_ref}  |  Result 总数: {stats.total_result}",
                f"*  匹配数:     {stats.matched_count}  |  断句: {stats.break_count}  |  连句: {stats.join_count}",
                f"",
                f"*  总检测率:   {stats.detection_rate * 100:6.2f}%  ({stats.total_ref - stats.fr_count}/{stats.total_ref})",
                f"*  FA 误检率:  {stats.fa_rate        * 100:6.2f}%  ({stats.fa_count}/{stats.total_result})",
                f"*  FR 漏检率:  {stats.fr_rate        * 100:6.2f}%  ({stats.fr_count}/{stats.total_ref})",
                f"",
                f"*  详细报告: {output_path}",
            ]
        return self.format_summary_box("VAD_ACCURACY", "\n".join(content_lines), width=90)


# ─── CLI 入口 ─────────────────────────────────────────────────────────────────

def main() -> None:
    """
    命令行入口函数

    使用示例:
        python3 vad_accuracy.py -r result.txt -ref ref.txt
        python3 vad_accuracy.py -r result.txt -ref ref.txt -o report.xlsx
    """
    import argparse

    parser = argparse.ArgumentParser(
        description='VAD FA/FR 准确率计算工具',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
示例:
  python3 %(prog)s -r result.txt -ref ref.txt
  python3 %(prog)s -r result.txt -ref ref.txt -o report.xlsx

输入文件格式（ref/result 相同格式，空格分隔，无表头）：
  音频路径  00h_S  开始时间(s)  结束时间(s)

输出:
  - 终端打印 Summary 统计报告
  - 生成 Excel 详细报告（5 个 Sheet：概要/FA详情/FR详情/Signal/Ref）
  - 生成同路径下的 HTML 报告
        '''
    )

    parser.add_argument('-r', '--result',    required=True,  help='VAD 结果文件路径')
    parser.add_argument('-ref', '--reference', required=True, help='参考 VAD 文件路径')
    parser.add_argument('-o', '--output',    default=None,   help='Excel 报告输出路径')

    args = parser.parse_args()

    params = [f"result={args.result}", f"ref={args.reference}"]
    if args.output:
        params.append(f"output={args.output}")

    handler = VADAccuracyHandler(config=None)
    try:
        summary = handler.execute(params)
        print(summary)
    except FileNotFoundError as e:
        print(f"错误: {e}"); exit(1)
    except ValueError as e:
        print(f"参数错误: {e}"); exit(1)
    except Exception as e:
        print(f"执行失败: {e}")
        import traceback; traceback.print_exc(); exit(1)


if __name__ == '__main__':
    main()
