#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# @Time    : 2026/03/04
# @Author  : huidong.bai
# @File    : vad_precision.py
# @Software: PyCharm
# @Mail    : MasterBai2018@outlook.com
"""
VAD_PRECISION 指令实现

用于分析 VAD（语音活动检测）的时间边界精度，计算：
- start 时间误差（识别起始时间 - 参考起始时间）
- end 时间误差（识别结束时间 - 参考结束时间）
- 60 / 70 / 80 / 90 百分位容忍值

匹配规则：贪心算法，优先选取重叠比例最大的候选段，
最低重叠比例阈值为 30%（= overlap / min(ref_dur, result_dur)）。

输入文件格式（ref 和 result 相同格式，空格分隔，无表头）：
    音频路径  00h_S  开始时间(s)  结束时间(s)

使用示例:
    [TSR]VAD_PRECISION result=vad.txt ref=ref.txt [output=report.xlsx]
"""
import os
import sys
import statistics
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Tuple, Optional
from loguru import logger

try:
    from . import register_tsr_command
    from .base import BaseTSRHandler
except ImportError:
    _current_dir = os.path.dirname(os.path.abspath(__file__))
    _project_root = os.path.abspath(os.path.join(_current_dir, '..', '..', '..', '..', '..'))
    if _project_root not in sys.path:
        sys.path.insert(0, _project_root)
    from src.testsuite.NANO.tools.tsr import register_tsr_command
    from src.testsuite.NANO.tools.tsr.base import BaseTSRHandler


# ─── 常量 ──────────────────────────────────────────────────────────────────────

MIN_OVERLAP_RATIO  = 0.3          # 最低重叠比例阈值（固定 30%）
PERCENTILE_LEVELS  = [60, 70, 80, 90]


# ─── 数据结构 ──────────────────────────────────────────────────────────────────

class VADSegment:
    """VAD 时间段（含重叠比例计算）"""

    def __init__(self, audio_path: str, start: float, end: float, line_num: int = 0):
        self.audio_path = audio_path
        self.start      = start
        self.end        = end
        self.line_num   = line_num

    def duration(self) -> float:
        return self.end - self.start

    def has_overlap(self, other: 'VADSegment') -> bool:
        return not (other.start >= self.end or other.end <= self.start)

    def overlap_ratio(self, other: 'VADSegment') -> float:
        """
        重叠比例 = 重叠时长 / min(自身时长, 对方时长)
        用于与 MIN_OVERLAP_RATIO 阈值对比
        """
        overlap_start = max(self.start, other.start)
        overlap_end   = min(self.end,   other.end)
        if overlap_start >= overlap_end:
            return 0.0
        overlap_dur = overlap_end - overlap_start
        min_dur     = min(self.duration(), other.duration())
        return overlap_dur / min_dur if min_dur > 0 else 0.0

    def __repr__(self) -> str:
        return f"VADSegment({self.audio_path}, {self.start:.3f}, {self.end:.3f})"


@dataclass
class MatchRecord:
    """一对匹配的 ref/result 及其误差"""
    audio_path:   str
    ref_start:    float
    ref_end:      float
    result_start: float
    result_end:   float
    start_error:  float   # result_start - ref_start（有正负方向）
    end_error:    float   # result_end   - ref_end（有正负方向）

    @property
    def abs_start_error(self) -> float:
        return abs(self.start_error)

    @property
    def abs_end_error(self) -> float:
        return abs(self.end_error)


@dataclass
class VADPrecisionStats:
    """VAD 时间精度汇总统计"""
    total_ref:        int = 0
    total_result:     int = 0
    matched_count:    int = 0
    unmatched_count:  int = 0   # ref 中未匹配的数量

    match_records:    List[MatchRecord]   = field(default_factory=list)

    # 误差列表（绝对值，用于百分位计算）
    abs_start_errors: List[float]         = field(default_factory=list)
    abs_end_errors:   List[float]         = field(default_factory=list)

    # 百分位统计（计算后填入）
    start_percentiles: Dict[int, float]   = field(default_factory=dict)
    end_percentiles:   Dict[int, float]   = field(default_factory=dict)
    is_fa_only:        bool               = False   # ref=NULL 时为 True

    @property
    def avg_start_error(self) -> float:
        return statistics.mean(self.abs_start_errors) if self.abs_start_errors else 0.0

    @property
    def avg_end_error(self) -> float:
        return statistics.mean(self.abs_end_errors) if self.abs_end_errors else 0.0

    @property
    def match_rate(self) -> float:
        return self.matched_count / self.total_ref if self.total_ref > 0 else 0.0


# ─── 文件解析 ──────────────────────────────────────────────────────────────────

def _parse_vad_file(file_path: str) -> List[VADSegment]:
    """
    解析 VAD 文件（ref 和 result 使用相同格式）
    格式：音频路径  00h_S  开始时间(s)  结束时间(s)
    """
    segments: List[VADSegment] = []
    with open(file_path, 'r', encoding='utf-8') as f:
        for line_num, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            parts = line.split()
            if len(parts) < 4:
                logger.warning(f"VAD 文件第 {line_num} 行格式不正确，跳过：{line}")
                continue
            try:
                audio_path = parts[0]
                start_time = float(parts[2])
                end_time   = float(parts[3])
                if start_time >= end_time:
                    logger.warning(f"VAD 文件第 {line_num} 行 start >= end，跳过")
                    continue
                segments.append(VADSegment(audio_path, start_time, end_time, line_num))
            except (ValueError, IndexError) as e:
                logger.warning(f"VAD 文件第 {line_num} 行解析失败，跳过：{e}")
    return segments


def _group_by_audio(segments: List[VADSegment]) -> Dict[str, List[VADSegment]]:
    """按音频路径分组并按开始时间排序"""
    grouped: Dict[str, List[VADSegment]] = defaultdict(list)
    for seg in segments:
        grouped[seg.audio_path].append(seg)
    for segs in grouped.values():
        segs.sort(key=lambda s: (s.start, s.end))
    return grouped


# ─── 贪心匹配算法 ──────────────────────────────────────────────────────────────

def _match_segments(
    ref_segs:    List[VADSegment],
    result_segs: List[VADSegment],
) -> Tuple[List[Tuple[VADSegment, VADSegment]], int]:
    """
    贪心匹配：对每个 ref 段，找重叠比例最大且 >= 30% 的 result 段。
    每个 result 段只能被匹配一次。

    Returns:
        matches:          匹配对列表
        unmatched_count:  未匹配的 ref 段数量
    """
    used_result_indices: set = set()
    matches: List[Tuple[VADSegment, VADSegment]] = []
    unmatched_count = 0

    for ref_seg in sorted(ref_segs, key=lambda s: s.start):
        best_idx     = -1
        best_ratio   = 0.0

        for idx, result_seg in enumerate(result_segs):
            if idx in used_result_indices:
                continue
            ratio = ref_seg.overlap_ratio(result_seg)
            if ratio >= MIN_OVERLAP_RATIO and ratio > best_ratio:
                best_ratio = ratio
                best_idx   = idx

        if best_idx >= 0:
            matches.append((ref_seg, result_segs[best_idx]))
            used_result_indices.add(best_idx)
        else:
            unmatched_count += 1

    return matches, unmatched_count


# ─── 百分位计算 ────────────────────────────────────────────────────────────────

def _calc_percentiles(values: List[float], levels: List[int]) -> Dict[int, float]:
    """线性插值计算百分位数"""
    if not values:
        return {p: 0.0 for p in levels}
    sv = sorted(values)
    result: Dict[int, float] = {}
    for p in levels:
        idx   = (p / 100.0) * (len(sv) - 1)
        lo    = int(idx)
        hi    = min(lo + 1, len(sv) - 1)
        w     = idx - lo
        result[p] = sv[lo] * (1 - w) + sv[hi] * w
    return result


# ─── 统计汇总 ──────────────────────────────────────────────────────────────────

def _calculate_stats(ref_file: str, result_file: str) -> VADPrecisionStats:
    """读取文件、按音频分组、贪心匹配、汇总误差统计"""
    result_segments = _parse_vad_file(result_file)
    logger.info(f"result 文件共 {len(result_segments)} 条记录")

    # ref=NULL 表示 FA-only 场景，无参考文件，直接返回空精度统计
    is_fa_only = ref_file.upper() == 'NULL'
    if is_fa_only:
        logger.info("FA-only 模式：ref=NULL，时间精度分析已跳过")
        return VADPrecisionStats(
            total_ref=0,
            total_result=len(result_segments),
            is_fa_only=True,
        )

    ref_segments = _parse_vad_file(ref_file)
    logger.info(f"ref 文件共 {len(ref_segments)} 条记录")

    ref_grouped    = _group_by_audio(ref_segments)
    result_grouped = _group_by_audio(result_segments)
    all_audio_paths = set(ref_grouped.keys()) | set(result_grouped.keys())

    stats = VADPrecisionStats(
        total_ref=len(ref_segments),
        total_result=len(result_segments),
    )

    for audio_path in sorted(all_audio_paths):
        ref_segs    = ref_grouped.get(audio_path, [])
        result_segs = result_grouped.get(audio_path, [])

        if not ref_segs:
            continue

        matches, unmatched = _match_segments(ref_segs, result_segs)
        stats.matched_count   += len(matches)
        stats.unmatched_count += unmatched

        for ref_seg, result_seg in matches:
            start_err = result_seg.start - ref_seg.start
            end_err   = result_seg.end   - ref_seg.end
            rec = MatchRecord(
                audio_path=audio_path,
                ref_start=ref_seg.start,
                ref_end=ref_seg.end,
                result_start=result_seg.start,
                result_end=result_seg.end,
                start_error=start_err,
                end_error=end_err,
            )
            stats.match_records.append(rec)
            stats.abs_start_errors.append(abs(start_err))
            stats.abs_end_errors.append(abs(end_err))

    # 计算百分位
    stats.start_percentiles = _calc_percentiles(stats.abs_start_errors, PERCENTILE_LEVELS)
    stats.end_percentiles   = _calc_percentiles(stats.abs_end_errors,   PERCENTILE_LEVELS)

    return stats


# ─── TSR 指令处理器 ───────────────────────────────────────────────────────────

@register_tsr_command("VAD_PRECISION")
class VADPrecisionHandler(BaseTSRHandler):
    """
    VAD 时间边界精度分析处理器

    指令格式:
        [TSR]VAD_PRECISION result=<result_txt> ref=<ref_txt> [output=<report_xlsx>]

    参数说明:
        result : VAD 识别结果文件路径（必需）
        ref    : 参考 VAD 文件路径（必需）
        output : Excel 报告输出路径（可选）

    输入文件格式（ref/result 相同）：
        音频路径  00h_S  开始时间(s)  结束时间(s)
    """

    COMMAND_NAME = "VAD_PRECISION"

    def execute(self, params: list, output_report: str = None) -> str:
        """执行 VAD 时间精度分析"""
        result_file, ref_file, xlsx_path, output_was_specified = self._parse_params(
            params, output_report
        )

        if not os.path.exists(result_file):
            raise FileNotFoundError(f"VAD 结果文件不存在: {result_file}")
        # ref="NULL" 为 FA-only 模式，无参考文件时跳过精度分析；否则文件必须存在
        if ref_file.upper() != 'NULL' and not os.path.exists(ref_file):
            raise FileNotFoundError(f"参考 VAD 文件不存在: {ref_file}")

        logger.info(f"VAD_PRECISION: result={result_file}, ref={ref_file}")

        stats = _calculate_stats(ref_file, result_file)

        # 生成 Excel 报告
        self._generate_excel_report(stats, ref_file, result_file, xlsx_path)

        # 生成 HTML 报告
        html_content  = self._build_precision_html(stats)
        html_filename = "vad_precision.html"
        self.save_html_report(html_content, html_filename, "VAD 时间精度报告", "vad_precision")

        # 本地副本
        local_html = os.path.splitext(xlsx_path)[0] + '.html'
        os.makedirs(os.path.dirname(local_html) or '.', exist_ok=True)
        with open(local_html, 'w', encoding='utf-8') as f:
            f.write(html_content)
        logger.info(f"本地 HTML 报告已生成: {local_html}")

        return self._generate_summary(stats, xlsx_path)

    # ── 参数解析 ──────────────────────────────────────────────────────────────

    def _parse_params(
        self, params: list, default_output: str = None
    ) -> Tuple[str, str, str, bool]:
        result_file = None
        ref_file    = None
        output_path = default_output

        for param in params:
            if   param.startswith("result="):
                result_file = param[7:]
            elif param.startswith("ref="):
                ref_file    = param[4:]
            elif param.startswith("output="):
                output_path = param[7:]

        if not result_file:
            raise ValueError("缺少必需参数: result=<result_txt>")
        if not ref_file:
            raise ValueError("缺少必需参数: ref=<ref_txt>")

        output_was_specified = output_path is not None
        if not output_path:
            suite_mango_dir = self._get_suite_mango_dir()
            output_path     = os.path.join(suite_mango_dir, "vad_precision.xlsx")

        return result_file, ref_file, output_path, output_was_specified

    # ── HTML 报告 ─────────────────────────────────────────────────────────────

    def generate_html_report(self, *args, **kwargs) -> Optional[str]:
        """实现抽象方法"""
        stats = args[0] if args else kwargs.get('stats')
        if stats is None:
            return None
        html = self._build_precision_html(stats)
        return self.save_html_report(html, "vad_precision.html", "VAD 时间精度报告", "vad_precision")

    def _build_precision_html(self, stats: VADPrecisionStats) -> str:
        """构建自包含 HTML"""

        # ── 汇总卡片 ──────────────────────────────────────────────────────────
        match_cls = ('card-success' if stats.match_rate >= 0.90 else
                     'card-warning' if stats.match_rate >= 0.70 else 'card-failure')
        cards_html = (
            '<div class="summary-cards">'
            f'<div class="card"><div class="card-value">{stats.total_ref}</div>'
            f'<div class="card-label">Ref 总数</div></div>'
            f'<div class="card {match_cls}"><div class="card-value">{stats.matched_count}</div>'
            f'<div class="card-label">匹配数</div></div>'
            f'<div class="card card-warning"><div class="card-value">{stats.unmatched_count}</div>'
            f'<div class="card-label">未匹配数</div></div>'
            f'<div class="card"><div class="card-value">{stats.avg_start_error * 1000:.1f} ms</div>'
            f'<div class="card-label">平均 Start 误差</div></div>'
            f'<div class="card"><div class="card-value">{stats.avg_end_error * 1000:.1f} ms</div>'
            f'<div class="card-label">平均 End 误差</div></div>'
            f'</div>'
        )

        # ── Section 1：百分位统计表 ────────────────────────────────────────────
        pct_rows = ''.join(
            f'<tr><td>{p}%</td>'
            f'<td>{stats.start_percentiles.get(p, 0) * 1000:.2f} ms</td>'
            f'<td>{stats.end_percentiles.get(p, 0) * 1000:.2f} ms</td></tr>'
            for p in PERCENTILE_LEVELS
        )
        pct_section = (
            '<div class="section">'
            '<div class="section-header">📊 误差百分位统计（绝对值）</div>'
            '<div class="section-body">'
            '<table class="report-table"><thead>'
            '<tr><th>百分位</th><th>Start 误差容忍值</th><th>End 误差容忍值</th></tr>'
            f'</thead><tbody>{pct_rows}</tbody></table>'
            '</div></div>'
        )

        # ── Section 2：匹配明细 Tab ────────────────────────────────────────────
        detail_rows = ''.join(
            f'<tr><td>{i}</td>'
            f'<td>{os.path.basename(r.audio_path)}</td>'
            f'<td>{r.ref_start:.3f}</td><td>{r.ref_end:.3f}</td>'
            f'<td>{r.result_start:.3f}</td><td>{r.result_end:.3f}</td>'
            f'<td><span class="tag-{"pass" if r.abs_start_error < 0.1 else "fail"}">'
            f'{r.start_error * 1000:+.1f} ms</span></td>'
            f'<td><span class="tag-{"pass" if r.abs_end_error < 0.1 else "fail"}">'
            f'{r.end_error * 1000:+.1f} ms</span></td>'
            f'</tr>'
            for i, r in enumerate(stats.match_records, 1)
        )
        detail_section = (
            '<div class="section">'
            '<div class="section-header">📋 匹配明细</div>'
            '<div class="section-body">'
            '<table class="report-table"><thead>'
            '<tr><th>#</th><th>音频</th>'
            '<th>Ref Start(s)</th><th>Ref End(s)</th>'
            '<th>Result Start(s)</th><th>Result End(s)</th>'
            '<th>Start 误差</th><th>End 误差</th></tr>'
            f'</thead><tbody>{detail_rows}</tbody></table>'
            '</div></div>'
        )

        return self.build_html_page(
            "VAD 时间精度报告", cards_html + pct_section + detail_section
        )

    # ── Excel 报告（2 个 Sheet）────────────────────────────────────────────────

    def _generate_excel_report(
        self,
        stats:       VADPrecisionStats,
        ref_file:    str,
        result_file: str,
        output_path: str,
    ) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("生成 Excel 报告需要安装 openpyxl：pip install openpyxl")

        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

        header_fill  = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font  = Font(bold=True, color="FFFFFF", size=11)
        title_font   = Font(bold=True, size=12)
        center_align = Alignment(horizontal='center', vertical='center')
        left_align   = Alignment(horizontal='left',   vertical='center')
        good_fill    = PatternFill(start_color="C5E0B4", end_color="C5E0B4", fill_type="solid")
        warn_fill    = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        def _write_header(ws, row, headers):
            for c, h in enumerate(headers, 1):
                cell = ws.cell(row=row, column=c, value=h)
                cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align

        # ── Sheet 1：汇总 ──────────────────────────────────────────────────────
        ws_sum = wb.create_sheet("汇总", 0)
        ws_sum['A1'] = 'VAD 时间精度分析结果'
        ws_sum['A1'].font = Font(bold=True, size=14)
        ws_sum.merge_cells('A1:B1')

        r = 3
        for label, value in [('参考文件：', ref_file), ('结果文件：', result_file),
                              ('重叠比例阈值：', f'{MIN_OVERLAP_RATIO * 100:.0f}%')]:
            ws_sum.cell(row=r, column=1, value=label)
            ws_sum.cell(row=r, column=2, value=value)
            r += 1

        r += 1
        ws_sum.cell(row=r, column=1, value='基本统计').font = title_font
        ws_sum.merge_cells(f'A{r}:B{r}')
        r += 1

        for label, value in [
            ('Ref 总段数',  stats.total_ref),
            ('Result 总段数', stats.total_result),
            ('成功匹配数',  stats.matched_count),
            ('未匹配数',    stats.unmatched_count),
            ('匹配率',      f'{stats.match_rate * 100:.2f}%'),
        ]:
            ws_sum.cell(row=r, column=1, value=label)
            ws_sum.cell(row=r, column=2, value=value)
            r += 1

        r += 1
        ws_sum.cell(row=r, column=1, value='误差统计（绝对值，单位：秒）').font = title_font
        ws_sum.merge_cells(f'A{r}:C{r}')
        r += 1
        _write_header(ws_sum, r, ['指标', 'Start 误差', 'End 误差'])
        r += 1

        for label, s_val, e_val in [
            ('平均误差',
             f'{stats.avg_start_error:.6f}', f'{stats.avg_end_error:.6f}'),
        ] + [
            (f'{p}% 百分位容忍值',
             f'{stats.start_percentiles.get(p, 0):.6f}',
             f'{stats.end_percentiles.get(p, 0):.6f}')
            for p in PERCENTILE_LEVELS
        ]:
            ws_sum.cell(row=r, column=1, value=label)
            ws_sum.cell(row=r, column=2, value=s_val)
            ws_sum.cell(row=r, column=3, value=e_val)
            r += 1

        ws_sum.column_dimensions['A'].width = 30
        ws_sum.column_dimensions['B'].width = 25
        ws_sum.column_dimensions['C'].width = 25

        # ── Sheet 2：误差明细 ──────────────────────────────────────────────────
        ws_detail = wb.create_sheet("误差明细", 1)
        detail_hdrs = [
            '音频路径',
            'Ref Start(s)', 'Ref End(s)',
            'Result Start(s)', 'Result End(s)',
            'Start 误差(s)', 'End 误差(s)',
            'Abs Start 误差(s)', 'Abs End 误差(s)',
        ]
        _write_header(ws_detail, 1, detail_hdrs)

        for ri, rec in enumerate(stats.match_records, 2):
            vals = [
                rec.audio_path,
                rec.ref_start, rec.ref_end,
                rec.result_start, rec.result_end,
                round(rec.start_error, 6), round(rec.end_error, 6),
                round(rec.abs_start_error, 6), round(rec.abs_end_error, 6),
            ]
            for ci, v in enumerate(vals, 1):
                cell           = ws_detail.cell(row=ri, column=ci, value=v)
                cell.alignment = center_align if ci > 1 else left_align
            # 误差较大时标红（abs > 0.1s）
            if rec.abs_start_error > 0.1:
                ws_detail.cell(row=ri, column=8).fill = warn_fill
            else:
                ws_detail.cell(row=ri, column=8).fill = good_fill
            if rec.abs_end_error > 0.1:
                ws_detail.cell(row=ri, column=9).fill = warn_fill
            else:
                ws_detail.cell(row=ri, column=9).fill = good_fill

        for ci, w in enumerate([60, 15, 15, 16, 16, 16, 16, 20, 20], 1):
            ws_detail.column_dimensions[get_column_letter(ci)].width = w
        ws_detail.freeze_panes = 'A2'

        parent = os.path.dirname(output_path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        wb.save(output_path)
        logger.info(f"Excel 报告已生成: {output_path}")

    # ── Summary 文本 ──────────────────────────────────────────────────────────

    def _generate_summary(self, stats: VADPrecisionStats, output_path: str) -> str:
        if stats.is_fa_only:
            # FA-only 模式：无参考文件，精度分析已跳过
            content_lines = [
                f"*  [FA-only 模式] 无参考文件，时间精度分析已跳过",
                f"",
                f"*  Result 总数:   {stats.total_result}",
                f"",
                f"*  详细报告: {output_path}",
            ]
            return self.format_summary_box("VAD_PRECISION", "\n".join(content_lines), width=90)

        pct_lines = [
            f"*  {p}% 容忍值:  Start={stats.start_percentiles.get(p, 0):.6f}s"
            f"  End={stats.end_percentiles.get(p, 0):.6f}s"
            for p in PERCENTILE_LEVELS
        ]
        content_lines = [
            f"*  Ref 总数: {stats.total_ref}  |  匹配: {stats.matched_count}"
            f"  |  未匹配: {stats.unmatched_count}"
            f"  |  匹配率: {stats.match_rate * 100:.2f}%",
            f"",
            f"*  平均 Start 误差: {stats.avg_start_error:.6f}s"
            f"  ({stats.avg_start_error * 1000:.1f} ms)",
            f"*  平均 End   误差: {stats.avg_end_error:.6f}s"
            f"  ({stats.avg_end_error * 1000:.1f} ms)",
            f"",
        ] + pct_lines + [
            f"",
            f"*  详细报告: {output_path}",
        ]
        return self.format_summary_box("VAD_PRECISION", "\n".join(content_lines), width=90)


# ─── CLI 入口 ─────────────────────────────────────────────────────────────────

def main() -> None:
    """
    命令行入口函数

    使用示例:
        python3 vad_precision.py -r result.txt -ref ref.txt
        python3 vad_precision.py -r result.txt -ref ref.txt -o report.xlsx
    """
    import argparse

    parser = argparse.ArgumentParser(
        description='VAD 时间边界精度分析工具',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=f'''
示例:
  python3 %(prog)s -r result.txt -ref ref.txt
  python3 %(prog)s -r result.txt -ref ref.txt -o report.xlsx

输入文件格式（ref/result 相同格式，空格分隔，无表头）：
  音频路径  00h_S  开始时间(s)  结束时间(s)

匹配规则：
  贪心算法，优先选取重叠比例最大的候选段，最低阈值 {MIN_OVERLAP_RATIO * 100:.0f}%

输出:
  - 终端打印 Summary（匹配数、平均误差、各百分位容忍值）
  - 生成 Excel 报告（2 个 Sheet：汇总/误差明细）
  - 生成同路径下的 HTML 报告
        '''
    )

    parser.add_argument('-r', '--result',    required=True,  help='VAD 结果文件路径')
    parser.add_argument('-ref', '--reference', required=True, help='参考 VAD 文件路径')
    parser.add_argument('-o', '--output',    default=None,   help='Excel 报告输出路径')

    args = parser.parse_args()

    params = [f"result={args.result}", f"ref={args.reference}"]
    if args.output:
        params.append(f"output={args.output}")

    handler = VADPrecisionHandler(config=None)
    try:
        summary = handler.execute(params)
        print(summary)
    except FileNotFoundError as e:
        print(f"错误: {e}"); exit(1)
    except ValueError as e:
        print(f"参数错误: {e}"); exit(1)
    except Exception as e:
        print(f"执行失败: {e}")
        import traceback; traceback.print_exc(); exit(1)


if __name__ == '__main__':
    main()
